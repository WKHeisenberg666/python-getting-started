"""The debt-overview Excel workbook, treated as the single source of truth.

Rather than mirroring a database into Excel (the earlier design), the dashboard reads this
file directly on every request and the ingestion pipeline writes/updates rows in it. A stable
"ID" column (the source document's checksum, or a random id for manually-added rows) lets rows
be found again for updates without relying on row position, which shifts as rows are added or
removed.
"""

import uuid
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from threading import Lock

import openpyxl

HEADER = [
    "ID",
    "Erkannt am",
    "Gläubiger",
    "Inkasso-Büro",
    "Kategorie",
    "Betrag",
    "Währung",
    "Bereits gezahlt",
    "Offener Betrag",
    "Fällig am",
    "Referenz",
    "Status",
    "Zusammenfassung",
    "Quelldokument",
]

# Internal (python-friendly) row keys, in the same order as HEADER.
FIELDS = [
    "id",
    "detected_at",
    "creditor_name",
    "collection_agency",
    "category",
    "amount",
    "currency",
    "paid_amount",
    "open_amount",
    "due_date",
    "reference",
    "status",
    "summary",
    "source_document",
]

STATUS_PROPOSAL = "Vorschlag (KI)"
STATUS_OPEN = "Offen"
STATUS_INSTALLMENT = "Ratenzahlung"
STATUS_PAID = "Bezahlt"
STATUS_SETTLED = "Beglichen"

# CSS-friendly class per status, for the dashboard badges.
STATUS_CSS_CLASS = {
    STATUS_PROPOSAL: "proposal",
    STATUS_OPEN: "open",
    STATUS_INSTALLMENT: "installment",
    STATUS_PAID: "paid",
    STATUS_SETTLED: "settled",
}

# A single process-wide lock: this is a single-user local tool, so a coarse lock that
# serializes read-modify-write access to the workbook file is simpler and safer than trying
# to do fine-grained per-row locking.
_lock = Lock()


@dataclass
class DebtRow:
    id: str
    detected_at: str = ""
    creditor_name: str = ""
    collection_agency: str = ""
    category: str = ""
    amount: Decimal = Decimal("0")
    currency: str = "EUR"
    paid_amount: Decimal = Decimal("0")
    open_amount: Decimal = Decimal("0")
    due_date: str = ""
    reference: str = ""
    status: str = STATUS_OPEN
    summary: str = ""
    source_document: str = ""
    extra: dict = field(default_factory=dict)

    def as_row(self) -> list:
        return [
            self.id,
            self.detected_at,
            self.creditor_name,
            self.collection_agency,
            self.category,
            float(self.amount or 0),
            self.currency,
            float(self.paid_amount or 0),
            float(self.open_amount if self.open_amount else self.amount or 0),
            self.due_date,
            self.reference,
            self.status,
            self.summary,
            self.source_document,
        ]


def _open_workbook(path: Path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        workbook = openpyxl.load_workbook(path)
        sheet = workbook.active
        return workbook, sheet
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Schulden"
    sheet.append(HEADER)
    return workbook, sheet


def _row_to_dict(row_values) -> dict:
    data = dict(zip(FIELDS, row_values))
    data["amount"] = Decimal(str(data["amount"] or 0))
    data["paid_amount"] = Decimal(str(data["paid_amount"] or 0))
    data["open_amount"] = Decimal(str(data["open_amount"] or 0))
    data["status_class"] = STATUS_CSS_CLASS.get(data["status"], "open")
    return data


def read_rows(path) -> list[dict]:
    with _lock:
        _, sheet = _open_workbook(path)
        rows = []
        for excel_row in sheet.iter_rows(min_row=2, values_only=True):
            if excel_row[0] is None:
                continue
            rows.append(_row_to_dict(excel_row))
        return rows


def _find_row_index(sheet, row_id: str) -> int | None:
    for idx, cell in enumerate(sheet["A"], start=1):
        if idx == 1:
            continue
        if cell.value == row_id:
            return idx
    return None


def new_id() -> str:
    return uuid.uuid4().hex


def upsert_row(path, debt_row: DebtRow) -> None:
    """Inserts a new row, or overwrites an existing one with the same ID."""
    with _lock:
        workbook, sheet = _open_workbook(path)
        existing_idx = _find_row_index(sheet, debt_row.id)
        values = debt_row.as_row()
        if existing_idx:
            for col, value in enumerate(values, start=1):
                sheet.cell(row=existing_idx, column=col, value=value)
        else:
            sheet.append(values)
        workbook.save(path)


def update_row(path, row_id: str, **changes) -> bool:
    """Updates only the given fields (by their FIELDS name) on the matching row."""
    with _lock:
        workbook, sheet = _open_workbook(path)
        idx = _find_row_index(sheet, row_id)
        if idx is None:
            return False
        for field_name, value in changes.items():
            col = FIELDS.index(field_name) + 1
            sheet.cell(row=idx, column=col, value=value)
        workbook.save(path)
        return True


def delete_row(path, row_id: str) -> bool:
    with _lock:
        workbook, sheet = _open_workbook(path)
        idx = _find_row_index(sheet, row_id)
        if idx is None:
            return False
        sheet.delete_rows(idx)
        workbook.save(path)
        return True


def today_str() -> str:
    return date.today().strftime("%d.%m.%Y")


def format_date(value: date | None) -> str:
    return value.strftime("%d.%m.%Y") if value else ""


def parse_date(value: str) -> date | None:
    if not value:
        return None
    try:
        return datetime.strptime(value, "%d.%m.%Y").date()
    except ValueError:
        return None
